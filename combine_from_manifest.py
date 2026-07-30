import csv
import os
import sys
from pathlib import Path

try:
    import tkinter as tk
    from tkinter import filedialog, messagebox
except Exception:
    tk = None
    filedialog = None
    messagebox = None

# Prefer local vendored packages if present.
_script_dir = Path(__file__).resolve().parent
for _vendor in (_script_dir / "_vendor", Path.cwd() / "_vendor"):
    if _vendor.exists():
        sys.path.insert(0, str(_vendor))

try:
    from pypdf import PdfMerger
except ImportError:
    try:
        from PyPDF2 import PdfMerger
    except ImportError as exc:
        raise SystemExit(
            "Missing PDF library. Install one of:\n"
            "  pip install pypdf\n"
            "or\n"
            "  pip install PyPDF2\n\n"
            "Tip: this script also supports local packages in a '_vendor' folder "
            "next to combine_from_manifest.py."
        ) from exc

try:
    from openpyxl import load_workbook
except Exception:
    load_workbook = None


def _pick_folder_with_dialog(initial_dir: Path) -> Path | None:
    if tk is None or filedialog is None:
        return None

    root = tk.Tk()
    root.withdraw()
    root.attributes("-topmost", True)
    selected = filedialog.askdirectory(
        title="Select PDF_OUTPUT folder",
        initialdir=str(initial_dir),
    )
    root.destroy()

    if not selected:
        return None
    return Path(selected)


def _show_message(title: str, message: str, is_error: bool = False) -> None:
    if messagebox is None:
        return
    root = tk.Tk()
    root.withdraw()
    root.attributes("-topmost", True)
    if is_error:
        messagebox.showerror(title, message)
    else:
        messagebox.showinfo(title, message)
    root.destroy()


def resolve_pdf_output_folder() -> Path:
    default_folder = Path.cwd() / "PDF_OUTPUT"
    env_folder = os.environ.get("PDF_OUTPUT_FOLDER")
    if env_folder:
        return Path(env_folder)

    if len(sys.argv) > 1 and sys.argv[1].strip():
        return Path(sys.argv[1].strip())

    dialog_choice = _pick_folder_with_dialog(default_folder)
    if dialog_choice is not None:
        return dialog_choice

    # Prompt only in interactive terminal sessions.
    if sys.stdin.isatty():
        answer = input(
            f"Enter PDF_OUTPUT folder path (blank for default: {default_folder}): "
        ).strip()
        if answer:
            return Path(answer)

    return default_folder


def should_continue(manifest_file: Path, ordered_rows: list[dict[str, str]]) -> bool:
    print(f"About to merge {len(ordered_rows)} PDFs from: {manifest_file}")
    for row in ordered_rows:
        order = row.get("Order", "?")
        pdf_path = Path(row["ExportedPdf"])
        print(f"  {order}: {pdf_path.name}")

    if not sys.stdin.isatty():
        return True

    answer = input("Proceed with merge? [Y/n]: ").strip().lower()
    return answer in ("", "y", "yes")


def load_manifest_rows(manifest_file: Path) -> list[dict[str, str]]:
    if manifest_file.suffix.lower() == ".xlsx":
        if load_workbook is None:
            raise RuntimeError(
                "Manifest is XLSX but openpyxl is not installed. "
                "Install with: pip install openpyxl"
            )
        workbook = load_workbook(filename=str(manifest_file), data_only=True)
        sheet = workbook.active
        rows = list(sheet.iter_rows(values_only=True))
        if not rows:
            return []
        headers = [str(c).strip() if c is not None else "" for c in rows[0]]
        if "ExportedPdf" not in headers:
            raise ValueError("Manifest must contain an 'ExportedPdf' column")
        items: list[dict[str, str]] = []
        for values in rows[1:]:
            row_dict: dict[str, str] = {}
            for idx, header in enumerate(headers):
                if not header:
                    continue
                value = values[idx] if idx < len(values) else ""
                row_dict[header] = "" if value is None else str(value)
            items.append(row_dict)
        return items

    with manifest_file.open("r", newline="", encoding="utf-8") as f:
        reader = csv.DictReader(f)
        if "ExportedPdf" not in (reader.fieldnames or []):
            raise ValueError("Manifest must contain an 'ExportedPdf' column")
        return list(reader)


def build_merge_rows(all_rows: list[dict[str, str]]) -> list[dict[str, str]]:
    ordered_rows: list[dict[str, str]] = []
    for row in all_rows:
        exported_pdf = (row.get("ExportedPdf") or "").strip()
        if not exported_pdf:
            order = row.get("Order", "?")
            status = row.get("Status", "Missing")
            model = row.get("ModelPath", "")
            print(f"Warning: skipping row {order} ({status}) with no ExportedPdf: {model}")
            continue

        pdf_path = Path(exported_pdf)
        if not pdf_path.exists():
            print(f"Warning: missing PDF from manifest: {pdf_path}")
            continue

        row["ExportedPdf"] = str(pdf_path)
        ordered_rows.append(row)

    return ordered_rows


def main() -> None:
    gui_mode = not sys.stdin.isatty()
    try:
        pdf_output_folder = resolve_pdf_output_folder()
        manifest_xlsx_candidates = [
            pdf_output_folder / "PDF_BOM_ORDER.xlsx",
            pdf_output_folder / "pdf_order_manifest.xlsx",
        ]
        manifest_csv_candidates = [
            pdf_output_folder / "PDF_BOM_ORDER.csv",
            pdf_output_folder / "pdf_order_manifest.csv",
        ]
        manifest_file = None
        for candidate in manifest_xlsx_candidates:
            if candidate.exists():
                manifest_file = candidate
                break

        # If XLSX is present but openpyxl is unavailable, gracefully fallback to CSV.
        if (
            manifest_file is not None
            and manifest_file.suffix.lower() == ".xlsx"
            and load_workbook is None
        ):
            csv_fallback = next((p for p in manifest_csv_candidates if p.exists()), None)
            if csv_fallback is not None:
                print(
                    "openpyxl not found; falling back to CSV manifest: "
                    f"{csv_fallback}"
                )
                manifest_file = csv_fallback

        if manifest_file is None:
            manifest_file = next((p for p in manifest_csv_candidates if p.exists()), manifest_xlsx_candidates[0])
        combined_output_file = pdf_output_folder / "Combined_BOM.pdf"

        if not manifest_file.exists():
            raise FileNotFoundError(f"Manifest not found: {manifest_file}")

        all_rows = load_manifest_rows(manifest_file)
        ordered_rows = build_merge_rows(all_rows)

        if not ordered_rows:
            raise RuntimeError("No PDFs found to merge from manifest")

        if not should_continue(manifest_file, ordered_rows):
            print("Merge cancelled by user.")
            return

        merger = PdfMerger()
        try:
            for row in ordered_rows:
                merger.append(row["ExportedPdf"])
            merger.write(str(combined_output_file))
        finally:
            merger.close()

        success_message = f"Merged {len(ordered_rows)} PDFs into:\n{combined_output_file}"
        print(success_message)
        if gui_mode:
            _show_message("BOM Merge Complete", success_message)

    except Exception as exc:
        if gui_mode:
            _show_message("BOM Merge Failed", str(exc), is_error=True)
        raise


if __name__ == "__main__":
    main()
